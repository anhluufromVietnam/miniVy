#!/usr/bin/env python3
"""
Local Full-Context API: Ollama + Qwen3-0.6B

Tự động đọc toàn bộ tài liệu trong thư mục docs/:
- .pdf
- .docx
- .txt
- .md
- .csv
- .json

Không dùng RAG / BGE / FAISS.
Mỗi request sẽ đọc lại toàn bộ tài liệu và đưa vào context của Qwen.

Install:
    pip install fastapi uvicorn requests pypdf python-docx

Run:
    ollama pull qwen3:0.6b
    python app_full_context.py
"""

from __future__ import annotations

import json
import os
from pathlib import Path

import requests
from docx import Document
from fastapi import FastAPI, HTTPException
from pydantic import BaseModel, Field
from pypdf import PdfReader


# =========================
# Configuration
# =========================

BASE_DIR = Path(__file__).resolve().parent
DOCS_DIR = BASE_DIR / "docs"
DOCS_DIR.mkdir(exist_ok=True)

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://127.0.0.1:11434")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "qwen3:0.6b")
PORT = int(os.getenv("PORT", "8000"))

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".json",
}

app = FastAPI(
    title="Local Qwen3 Full Context API",
    version="2.0.0",
)


# =========================
# Request model
# =========================

class ChatRequest(BaseModel):
    message: str = Field(min_length=1)
    temperature: float = 0.2
    max_tokens: int = 512


# =========================
# Document readers
# =========================

def read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages: list[str] = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
            if text.strip():
                pages.append(f"[PAGE {page_number}]\n{text}")
        except Exception as exc:
            print(f"[PDF] Failed page {page_number} in {path}: {exc}")

    return "\n\n".join(pages)


def read_docx(path: Path) -> str:
    doc = Document(str(path))
    parts: list[str] = []

    # Paragraphs
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    # Tables
    for table_index, table in enumerate(doc.tables, start=1):
        parts.append(f"[TABLE {table_index}]")

        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return read_pdf(path)

    if suffix == ".docx":
        return read_docx(path)

    if suffix in {".txt", ".md", ".csv"}:
        return read_text_file(path)

    if suffix == ".json":
        raw = read_text_file(path)

        try:
            data = json.loads(raw)
            return json.dumps(
                data,
                ensure_ascii=False,
                indent=2,
            )
        except json.JSONDecodeError:
            return raw

    return ""


def load_all_documents() -> tuple[str, list[str]]:
    """
    Đọc toàn bộ PDF/DOCX/TXT/MD/CSV/JSON trong docs/.
    Không chunk, không embedding, không retrieval.
    """

    documents: list[str] = []
    file_names: list[str] = []

    for path in sorted(DOCS_DIR.rglob("*")):
        if not path.is_file():
            continue

        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        try:
            text = read_document(path).strip()

            if not text:
                print(f"[DOCS] Empty or unreadable: {path}")
                continue

            relative_path = str(path.relative_to(BASE_DIR))

            documents.append(
                f"===== DOCUMENT: {relative_path} =====\n"
                f"{text}"
            )
            file_names.append(relative_path)

            print(
                f"[DOCS] Loaded: {relative_path} "
                f"({len(text):,} characters)"
            )

        except Exception as exc:
            print(f"[DOCS] Failed to read {path}: {exc}")

    return "\n\n".join(documents), file_names


def get_document_context() -> tuple[str, list[str]]:
    # Đọc lại mỗi request, không cần restart khi thay tài liệu.
    return load_all_documents()


# =========================
# Ollama
# =========================

def ollama_chat(
    messages: list[dict[str, str]],
    temperature: float = 0.2,
    max_tokens: int = 512,
) -> str:

    payload = {
        "model": OLLAMA_MODEL,
        "messages": messages,
        "stream": False,
        "options": {
            "temperature": temperature,
            "num_predict": max_tokens,
        },
    }

    try:
        response = requests.post(
            f"{OLLAMA_URL}/api/chat",
            json=payload,
            timeout=300,
        )
        response.raise_for_status()

    except requests.RequestException as exc:
        raise HTTPException(
            status_code=503,
            detail=(
                "Ollama is not reachable. "
                f"Expected {OLLAMA_URL}. "
                f"Error: {exc}"
            ),
        )

    data = response.json()
    return data.get("message", {}).get("content", "")


# =========================
# Routes
# =========================

@app.get("/health")
def health():
    ollama_ok = False

    try:
        response = requests.get(
            f"{OLLAMA_URL}/api/tags",
            timeout=3,
        )
        ollama_ok = response.ok
    except requests.RequestException:
        pass

    context, files = get_document_context()

    return {
        "status": "ok",
        "ollama": ollama_ok,
        "ollama_url": OLLAMA_URL,
        "ollama_model": OLLAMA_MODEL,
        "documents": files,
        "supported_extensions": sorted(SUPPORTED_EXTENSIONS),
        "context_characters": len(context),
        "mode": "full-context",
    }


@app.get("/document")
def document():
    context, files = get_document_context()

    return {
        "files": files,
        "characters": len(context),
        "context": context,
    }


@app.post("/chat")
def chat(req: ChatRequest):
    context, files = get_document_context()

    if not context:
        context = "(Không có tài liệu được tìm thấy trong thư mục docs.)"

    system_prompt = """Bạn là trợ lý AI chạy hoàn toàn offline.

Bạn được cung cấp TOÀN BỘ nội dung các tài liệu trong CONTEXT.

Quy tắc:
1. Ưu tiên thông tin trong CONTEXT.
2. Có thể tổng hợp thông tin từ nhiều tài liệu và nhiều vị trí.
3. Không tự bịa thông tin không có trong CONTEXT.
4. Nếu không tìm thấy thông tin, nói rõ:
   "Tôi không tìm thấy thông tin này trong tài liệu."
5. Trả lời bằng ngôn ngữ của người dùng.
6. Giữ câu trả lời rõ ràng, trực tiếp và có cấu trúc.
7. Với câu hỏi tiếp nối, dùng CONTEXT để hiểu lại nội dung.
"""

    user_prompt = f"""CONTEXT — TOÀN BỘ TÀI LIỆU:
{context}

========================

QUESTION:
{req.message}
"""

    answer = ollama_chat(
        [
            {
                "role": "system",
                "content": system_prompt,
            },
            {
                "role": "user",
                "content": user_prompt,
            },
        ],
        temperature=req.temperature,
        max_tokens=req.max_tokens,
    )

    return {
        "answer": answer,
        "model": OLLAMA_MODEL,
        "mode": "full-context",
        "documents": files,
        "context_characters": len(context),
    }


# =========================
# Startup
# =========================

@app.on_event("startup")
def startup():
    print("=" * 60)
    print("Local Qwen3 Full Context API")
    print("=" * 60)
    print(f"Ollama: {OLLAMA_URL}")
    print(f"LLM:    {OLLAMA_MODEL}")
    print(f"Docs:   {DOCS_DIR}")
    print(f"Types:  {', '.join(sorted(SUPPORTED_EXTENSIONS))}")
    print(f"API:    http://127.0.0.1:{PORT}")
    print("=" * 60)

    context, files = get_document_context()

    print(f"[DOCS] Files loaded: {len(files)}")
    print(f"[DOCS] Context size: {len(context):,} characters")


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        app,
        host="127.0.0.1",
        port=PORT,
    )
