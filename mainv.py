#!/usr/bin/env python3
"""
Local Voice Assistant API
FastAPI + Ollama + Qwen3

Features
- Reads documents from docs/
- Supports PDF DOCX TXT MD CSV JSON
- Full-context prompting
- Streaming Ollama responses
- Short English voice-conversation responses
"""

from __future__ import annotations
from fastapi.middleware.cors import CORSMiddleware
import json
import httpx
import os
from pathlib import Path

import requests
from docx import Document
from fastapi import FastAPI, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from pypdf import PdfReader


BASE_DIR = Path(__file__).resolve().parent
DOCS_DIR = BASE_DIR / "docs"
DOCS_DIR.mkdir(exist_ok=True)

OLLAMA_URL = os.getenv("OLLAMA_URL", "http://127.0.0.1:11434")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "qwen3:0.6b")
PORT = int(os.getenv("PORT", "8000"))

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".csv",
    ".json",
}


app = FastAPI(
    title="Local Qwen Voice Assistant API",
    version="3.0.0",
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)
class ChatRequest(BaseModel):
    message: str = Field(min_length=1)
    temperature: float = 0.2
    max_tokens: int = 2048


def read_text_file(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""

            if text.strip():
                pages.append(
                    f"[PAGE {page_number}]\n{text}"
                )

        except Exception as exc:
            print(
                f"[PDF] Failed page {page_number} "
                f"in {path}: {exc}"
            )

    return "\n\n".join(pages)


def read_docx(path: Path) -> str:
    doc = Document(str(path))
    parts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    for table_index, table in enumerate(
        doc.tables,
        start=1,
    ):
        parts.append(f"[TABLE {table_index}]")

        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def read_document(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return read_pdf(path)

    if suffix == ".docx":
        return read_docx(path)

    if suffix in {
        ".txt",
        ".md",
        ".csv",
    }:
        return read_text_file(path)

    if suffix == ".json":
        raw = read_text_file(path)

        try:
            data = json.loads(raw)

            return json.dumps(
                data,
                ensure_ascii=False,
                indent=2,
            )

        except json.JSONDecodeError:
            return raw

    return ""


def load_all_documents():
    documents = []
    file_names = []

    for path in sorted(DOCS_DIR.rglob("*")):

        if not path.is_file():
            continue

        if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        try:
            text = read_document(path).strip()

            if not text:
                continue

            relative_path = str(
                path.relative_to(BASE_DIR)
            )

            documents.append(
                f"===== DOCUMENT: {relative_path} =====\n"
                f"{text}"
            )

            file_names.append(relative_path)

        except Exception as exc:
            print(
                f"[DOCS] Failed to read "
                f"{path}: {exc}"
            )

    return "\n\n".join(documents), file_names


def get_document_context():
    return load_all_documents()


VOICE_SYSTEM_PROMPT = """
You are a friendly English voice conversation partner

Speak English by default

Keep every response very short

Usually use one short sentence

Use simple natural spoken English

Use contractions naturally

Sound like a real person having a casual conversation

Do not sound formal or academic

Do not give long explanations unless the user asks

Do not repeat the users question

Ask at most one simple follow up question

Do not use markdown

Do not use bullet points

Do not use emojis

Use normal punctuation internally so the streaming client can detect sentences

The client will remove punctuation before speech synthesis

When information is available in CONTEXT use it as the primary source

Do not invent facts that are not supported by CONTEXT

If the requested information is not in CONTEXT say

I cannot find that information in the available documents

Prioritize natural conversation over explanations
"""


def build_messages(message: str, context: str):

    if not context:
        context = (
            "(No documents were found in the docs folder)"
        )

    user_prompt = f"""
CONTEXT

{context}

END CONTEXT

USER

{message}
"""

    return [
        {
            "role": "system",
            "content": VOICE_SYSTEM_PROMPT,
        },
        {
            "role": "user",
            "content": user_prompt,
        },
    ]


def ollama_stream(
    messages,
    temperature=0.2,
    max_tokens=2048,
):
    payload = {
        "model": OLLAMA_MODEL,
        "messages": messages,
        "stream": True,
        "options": {
            "temperature": temperature,
            "num_predict": max_tokens,
        },
    }

    try:
        response = requests.post(
            f"{OLLAMA_URL}/api/chat",
            json=payload,
            stream=True,
            timeout=300,
        )

        response.raise_for_status()

    except requests.RequestException as exc:
        raise HTTPException(
            status_code=503,
            detail=(
                "Ollama is not reachable. "
                f"Expected {OLLAMA_URL}. "
                f"Error: {exc}"
            ),
        )

    for line in response.iter_lines():

        if not line:
            continue

        try:
            data = json.loads(line)

            content = (
                data
                .get("message", {})
                .get("content", "")
            )

            if content:
                yield content

            if data.get("done"):
                break

        except json.JSONDecodeError:
            continue


@app.get("/health")
def health():

    ollama_ok = False

    try:
        response = requests.get(
            f"{OLLAMA_URL}/api/tags",
            timeout=3,
        )

        ollama_ok = response.ok

    except requests.RequestException:
        pass

    context, files = get_document_context()

    return {
        "status": "ok",
        "ollama": ollama_ok,
        "ollama_url": OLLAMA_URL,
        "ollama_model": OLLAMA_MODEL,
        "documents": files,
        "context_characters": len(context),
        "mode": "voice-stream",
    }

@app.post("/chat/stream")
async def chat_stream(req: ChatRequest):
    # Đọc toàn bộ documents mỗi request
    # để luôn lấy nội dung mới nhất trong docs/
    context, files = get_document_context()

    messages = build_messages(
        req.message,
        context,
    )

    async def generate():
        payload = {
            "model": OLLAMA_MODEL,

            # Qwen3:
            # Chỉ stream câu trả lời, không stream thinking
            "think": False,

            "messages": messages,

            "stream": True,

            "options": {
                "temperature": req.temperature,
                "num_predict": req.max_tokens,
            },
        }

        print("=" * 60, flush=True)
        print("[STREAM] Request:", req.message, flush=True)
        print("[STREAM] Documents:", files, flush=True)
        print(
            "[STREAM] Context characters:",
            len(context),
            flush=True,
        )
        print(
            "[STREAM] Temperature:",
            req.temperature,
            flush=True,
        )
        print(
            "[STREAM] Max tokens:",
            req.max_tokens,
            flush=True,
        )
        print("=" * 60, flush=True)

        try:
            async with httpx.AsyncClient(
                timeout=None
            ) as client:

                async with client.stream(
                    "POST",
                    f"{OLLAMA_URL}/api/chat",
                    json=payload,
                ) as response:

                    print(
                        "[STREAM] Ollama status:",
                        response.status_code,
                        flush=True,
                    )

                    response.raise_for_status()

                    async for line in response.aiter_lines():

                        if not line:
                            continue

                        try:
                            data = json.loads(line)

                        except json.JSONDecodeError:
                            print(
                                "[STREAM] Invalid JSON:",
                                line,
                                flush=True,
                            )
                            continue

                        # Chỉ lấy content
                        # Không lấy thinking
                        content = (
                            data
                            .get("message", {})
                            .get("content", "")
                        )

                        if content:
                            print(
                                "[STREAM] TOKEN:",
                                repr(content),
                                flush=True,
                            )

                            yield content

                        if data.get("done"):
                            print(
                                "[STREAM] DONE:",
                                data.get("done_reason"),
                                flush=True,
                            )
                            break

        except httpx.HTTPError as exc:
            print(
                "[STREAM] HTTP ERROR:",
                exc,
                flush=True,
            )

        except Exception as exc:
            print(
                "[STREAM] ERROR:",
                exc,
                flush=True,
            )

    return StreamingResponse(
        generate(),
        media_type="text/plain; charset=utf-8",
        headers={
            "Cache-Control": "no-cache, no-transform",
            "X-Accel-Buffering": "no",
            "Connection": "keep-alive",
        },
    )

@app.post("/chat")
def chat(req: ChatRequest):

    context, files = get_document_context()

    messages = build_messages(
        req.message,
        context,
    )

    answer = "".join(
        ollama_stream(
            messages,
            temperature=req.temperature,
            max_tokens=min(req.max_tokens, 2048),
        )
    )

    return {
        "answer": answer,
        "model": OLLAMA_MODEL,
        "documents": files,
        "mode": "voice",
    }


@app.get("/document")
def document():

    context, files = get_document_context()

    return {
        "files": files,
        "characters": len(context),
        "context": context,
    }


@app.on_event("startup")
def startup():

    print("=" * 60)
    print("Local Qwen Voice Assistant API")
    print("=" * 60)
    print(f"Ollama: {OLLAMA_URL}")
    print(f"Model:  {OLLAMA_MODEL}")
    print(f"Docs:   {DOCS_DIR}")
    print(f"API:    http://127.0.0.1:{PORT}")
    print("=" * 60)


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(
        app,
        host="127.0.0.1",
        port=PORT,
    )
